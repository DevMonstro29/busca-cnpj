# -*- coding: utf-8 -*-
"""
Sistema de Busca CNPJ + Perfil da Empresa no Google
- Consulta dados do CNPJ na BrasilAPI (Receita Federal)
- Verifica Perfil da Empresa via Google Places API (New)
"""
import io
import re
import time
import difflib
import unicodedata

import requests
from flask import Flask, jsonify, render_template, request

app = Flask(__name__)

BRASILAPI_URL = "https://brasilapi.com.br/api/cnpj/v1/{cnpj}"
MINHARECEITA_URL = "https://minhareceita.org/{cnpj}"
PLACES_URL = "https://places.googleapis.com/v1/places:searchText"

CNPJ_REGEX = re.compile(r"\d{2}\.?\d{3}\.?\d{3}\s*/?\s*\d{4}-?\d{2}")

# ---------------------------------------------------------------------------
# Categorias de negócio permitidas (filtro do Perfil Google)
# ---------------------------------------------------------------------------
# A Places API não expõe os códigos GCID (3512/55/3477). O filtro casa o texto
# da categoria que o Google exibe (primaryTypeDisplayName) e os "types" técnicos.
# Cada grupo lista palavras-chave (já sem acento e minúsculas) que identificam
# a categoria desejada.
CATEGORIAS_PERMITIDAS = {
    # 3512 - Saúde e beleza > Mobilidade e acessibilidade > Scooters elétricos
    # + Scooters/bicicletas (mobilidade leve)
    "Scooters / bike / mobilidade": [
        "scooter",
        "patinete",
        "bicicleta",
        "bike",
        "triciclo",
        "mobilidade",
        "acessibilidade",
        "cadeira de rodas",
        "equipamento de mobilidade",
        "mobility",
        "bicycle",
    ],
    # 55 - Artes e entretenimento > Acessórios para instrumentos musicais
    "Acessórios para instrumentos musicais": [
        "instrumento musical",
        "instrumentos musicais",
        "acessorios para instrumentos",
        "loja de instrumentos",
        "musical instrument",
    ],
    # 3477 - Comercial e industrial > Medicina > Equipamento médico
    "Equipamento médico": [
        "equipamento medico",
        "equipamentos medicos",
        "produtos medicos",
        "material medico",
        "loja de produtos medicos",
        "ortopedic",
        "artigos medicos",
        "medical supply",
        "medical equipment",
    ],
    # Náutica (vendas relacionadas a embarcações/náutica)
    "Náutica": [
        "nautic",
        "barco",
        "embarca",
        "iate",
        "lancha",
        "marina",
        "motor de popa",
        "estaleiro",
        "boat",
        "marine",
        "veiculos recreativos",
    ],
}


# ---------------------------------------------------------------------------
# Utilidades de CNPJ
# ---------------------------------------------------------------------------
def somente_digitos(texto: str) -> str:
    return re.sub(r"\D", "", texto or "")


def cnpj_valido(cnpj: str) -> bool:
    """Valida os dígitos verificadores do CNPJ."""
    cnpj = somente_digitos(cnpj)
    if len(cnpj) != 14 or cnpj == cnpj[0] * 14:
        return False

    def calc_digito(base: str, pesos: list) -> int:
        soma = sum(int(d) * p for d, p in zip(base, pesos))
        resto = soma % 11
        return 0 if resto < 2 else 11 - resto

    pesos1 = [5, 4, 3, 2, 9, 8, 7, 6, 5, 4, 3, 2]
    pesos2 = [6] + pesos1
    d1 = calc_digito(cnpj[:12], pesos1)
    d2 = calc_digito(cnpj[:12] + str(d1), pesos2)
    return cnpj[12] == str(d1) and cnpj[13] == str(d2)


def formatar_cnpj(cnpj: str) -> str:
    c = somente_digitos(cnpj)
    if len(c) != 14:
        return cnpj
    return f"{c[:2]}.{c[2:5]}.{c[5:8]}/{c[8:12]}-{c[12:]}"


def extrair_cnpjs(texto: str) -> list:
    """Extrai CNPJs de um texto livre (com ou sem formatação)."""
    encontrados = []
    vistos = set()
    for m in CNPJ_REGEX.finditer(texto):
        c = somente_digitos(m.group())
        if len(c) == 14 and c not in vistos:
            vistos.add(c)
            encontrados.append(c)
    # fallback: linhas com 12 a 14 dígitos (Excel costuma cortar zeros à esquerda)
    for linha in texto.splitlines():
        c = somente_digitos(linha)
        if 12 <= len(c) <= 14:
            if len(c) < 14:
                c = c.zfill(14)
                if not cnpj_valido(c):
                    continue  # só aceita CNPJ completado se os dígitos verificadores baterem
            if c not in vistos:
                vistos.add(c)
                encontrados.append(c)
    return encontrados


# ---------------------------------------------------------------------------
# Consulta BrasilAPI (dados da Receita Federal)
# ---------------------------------------------------------------------------
def _montar_dados(d: dict) -> dict:
    return {
        "razao_social": d.get("razao_social") or "",
        "nome_fantasia": d.get("nome_fantasia") or "",
        "segmento": d.get("cnae_fiscal_descricao") or "",
        "municipio": d.get("municipio") or "",
        "uf": d.get("uf") or "",
        "situacao": d.get("descricao_situacao_cadastral") or "",
        "telefone": d.get("ddd_telefone_1") or "",
        "cep": str(d.get("cep") or ""),
        "logradouro": d.get("logradouro") or "",
        "bairro": d.get("bairro") or "",
    }


def consultar_receita(cnpj: str) -> dict:
    """Consulta dados do CNPJ. Tenta BrasilAPI e usa minhareceita.org como reserva."""
    erro_brasilapi = None
    try:
        r = requests.get(BRASILAPI_URL.format(cnpj=cnpj), timeout=20)
        if r.status_code == 200:
            return _montar_dados(r.json())
        if r.status_code == 404:
            return {"erro": "CNPJ não encontrado na Receita Federal"}
        erro_brasilapi = f"HTTP {r.status_code}"
    except requests.RequestException as e:
        erro_brasilapi = str(e)

    # Fallback: minhareceita.org
    try:
        time.sleep(0.5)
        r = requests.get(MINHARECEITA_URL.format(cnpj=cnpj), timeout=20)
        if r.status_code == 200:
            return _montar_dados(r.json())
        if r.status_code == 404:
            return {"erro": "CNPJ não encontrado na Receita Federal"}
        return {"erro": f"Erro nas consultas (BrasilAPI: {erro_brasilapi} / MinhaReceita: HTTP {r.status_code})"}
    except requests.RequestException as e:
        return {"erro": f"Falha de conexão (BrasilAPI: {erro_brasilapi} / MinhaReceita: {e})"}


# ---------------------------------------------------------------------------
# Consulta Google Places API (New) — Perfil da Empresa
# ---------------------------------------------------------------------------
# Sufixos societários que o Google geralmente não exibe no nome do perfil
SUFIXOS_RE = re.compile(
    r"\b(ltda|limitada|me|mei|epp|eireli|s\s*/?\s*a|sa|sociedade\s+anonima)\b",
    re.IGNORECASE,
)


def normalizar_nome(nome: str) -> str:
    """Remove acentos, pontuação e sufixos societários (LTDA, ME, EPP...)."""
    nome = unicodedata.normalize("NFKD", nome or "")
    nome = "".join(ch for ch in nome if not unicodedata.combining(ch))
    nome = nome.lower()
    nome = re.sub(r"[^a-z0-9\s]", " ", nome)
    nome = SUFIXOS_RE.sub(" ", nome)
    return re.sub(r"\s+", " ", nome).strip()


def similaridade(a: str, b: str) -> float:
    """Compara nomes normalizados: sequência de caracteres + tokens em comum."""
    a, b = normalizar_nome(a), normalizar_nome(b)
    if not a or not b:
        return 0.0
    seq = difflib.SequenceMatcher(None, a, b).ratio()
    ta, tb = set(a.split()), set(b.split())
    comuns = len(ta & tb)
    jaccard = comuns / len(ta | tb)
    contencao = comuns / min(len(ta), len(tb))  # nome curto contido no longo
    return max(seq, (jaccard + contencao) / 2)


def telefones_iguais(t1: str, t2: str) -> bool:
    """Compara telefones ignorando formatação e código do país."""
    d1, d2 = somente_digitos(t1), somente_digitos(t2)
    if len(d1) < 10 or len(d2) < 10:
        return False
    if d1.startswith("55") and len(d1) > 11:
        d1 = d1[2:]
    if d2.startswith("55") and len(d2) > 11:
        d2 = d2[2:]
    # mesmo DDD e mesmos 8 últimos dígitos (cobre celular com/sem o 9)
    return d1[:2] == d2[:2] and d1[-8:] == d2[-8:]


def categoria_do_place(p: dict) -> str:
    """Categoria que o Google exibe, sem acento/minúscula (para casamento)."""
    nome_tipo = (p.get("primaryTypeDisplayName") or {}).get("text", "")
    return normalizar_nome(nome_tipo)


def categoria_permitida(p: dict) -> str:
    """Retorna o nome do grupo permitido se o place se encaixar; senão, ''.

    Casa o texto da categoria exibida (primaryTypeDisplayName) e os 'types'
    técnicos contra as palavras-chave de CATEGORIAS_PERMITIDAS.
    """
    texto_categoria = categoria_do_place(p)
    types = " ".join(p.get("types") or [])
    primary_type = p.get("primaryType") or ""
    alvo = f"{texto_categoria} {types} {primary_type}".lower()

    for grupo, palavras in CATEGORIAS_PERMITIDAS.items():
        for palavra in palavras:
            if normalizar_nome(palavra) in alvo or palavra.lower() in alvo:
                return grupo
    return ""


def pontuar_place(p: dict, dados: dict) -> tuple:
    """Pontua um resultado do Google cruzando nome, telefone e endereço."""
    nome_google = (p.get("displayName") or {}).get("text", "")
    score_nome = max(
        similaridade(nome_google, dados.get("nome_fantasia", "")),
        similaridade(nome_google, dados.get("razao_social", "")),
    )
    score = score_nome
    sinais = []

    endereco = p.get("formattedAddress", "")
    tel_google = p.get("nationalPhoneNumber") or p.get("internationalPhoneNumber") or ""

    # telefone igual é o sinal mais forte
    if telefones_iguais(tel_google, dados.get("telefone", "")):
        score += 0.35
        sinais.append("telefone confere")

    # CEP do perfil bate com o da Receita
    cep = somente_digitos(dados.get("cep", ""))
    if cep and cep in somente_digitos(endereco):
        score += 0.25
        sinais.append("CEP confere")
    else:
        # rua: maioria das palavras do logradouro presente no endereço do Google
        rua = normalizar_nome(dados.get("logradouro", ""))
        end_norm = normalizar_nome(endereco)
        tokens_rua = [t for t in rua.split() if len(t) > 2]
        if tokens_rua and sum(t in end_norm for t in tokens_rua) >= max(1, len(tokens_rua) - 1):
            score += 0.15
            sinais.append("rua confere")

    # cidade divergente derruba a confiança (homônimo de outro município)
    municipio = normalizar_nome(dados.get("municipio", ""))
    if municipio and municipio not in normalizar_nome(endereco):
        score -= 0.25
        sinais.append("cidade divergente")

    return min(score, 1.0), score_nome, sinais


def consultar_google(api_key: str, dados: dict) -> dict:
    nome_busca = dados.get("nome_fantasia") or dados.get("razao_social")
    if not nome_busca:
        return {"perfil_google": "Não", "detalhe_google": "Sem nome para buscar"}

    query = " ".join(
        filter(None, [nome_busca, dados.get("bairro", ""), dados.get("municipio", ""), dados.get("uf", "")])
    )
    headers = {
        "Content-Type": "application/json",
        "X-Goog-Api-Key": api_key,
        "X-Goog-FieldMask": (
            "places.displayName,places.formattedAddress,places.rating,"
            "places.userRatingCount,places.googleMapsUri,places.businessStatus,"
            "places.nationalPhoneNumber,places.internationalPhoneNumber,"
            "places.websiteUri,"
            "places.types,places.primaryType,places.primaryTypeDisplayName"
        ),
    }
    body = {"textQuery": query, "languageCode": "pt-BR", "regionCode": "BR", "maxResultCount": 5}

    try:
        r = requests.post(PLACES_URL, json=body, headers=headers, timeout=20)
        if r.status_code == 403:
            return {"perfil_google": "Erro", "detalhe_google": "Chave de API inválida ou Places API (New) não habilitada"}
        if r.status_code != 200:
            return {"perfil_google": "Erro", "detalhe_google": f"Erro Google (HTTP {r.status_code})"}

        places = r.json().get("places", [])
        if not places:
            return {"perfil_google": "Não", "detalhe_google": "Nenhum perfil encontrado"}

        # mantém só os perfis cuja categoria é uma das permitidas
        permitidos = [(p, categoria_permitida(p)) for p in places]
        permitidos = [(p, g) for p, g in permitidos if g]
        if not permitidos:
            return {"perfil_google": "Não", "detalhe_google": "Fora das categorias permitidas"}

        # escolhe o resultado com melhor pontuação combinada
        melhor, melhor_score, melhor_nome, melhor_sinais, melhor_grupo = None, 0.0, 0.0, [], ""
        for p, grupo in permitidos:
            score, score_nome, sinais = pontuar_place(p, dados)
            if score > melhor_score:
                melhor, melhor_score, melhor_nome, melhor_sinais, melhor_grupo = p, score, score_nome, sinais, grupo

        if melhor is None or melhor_score < 0.40:
            return {"perfil_google": "Não", "detalhe_google": "Nenhum perfil correspondente encontrado"}

        nome_google = (melhor.get("displayName") or {}).get("text", "")
        site = melhor.get("websiteUri", "") or ""
        resultado = {
            "nome_google": nome_google,
            "endereco_google": melhor.get("formattedAddress", ""),
            "telefone_google": melhor.get("nationalPhoneNumber", ""),
            "avaliacao": melhor.get("rating", ""),
            "total_avaliacoes": melhor.get("userRatingCount", ""),
            "link_maps": melhor.get("googleMapsUri", ""),
            "categoria_google": (melhor.get("primaryTypeDisplayName") or {}).get("text", ""),
            "grupo_categoria": melhor_grupo,
            "site": site,
            "tem_site": bool(site.strip()),
        }

        detalhe = f"Nome {melhor_nome:.0%}"
        if melhor_sinais:
            detalhe += " · " + ", ".join(melhor_sinais)
        detalhe += " · SEM SITE" if not site.strip() else " · com site"

        if melhor.get("businessStatus") == "CLOSED_PERMANENTLY":
            resultado["perfil_google"] = "Possível"
            resultado["detalhe_google"] = detalhe + " · perfil marcado como fechado permanentemente"
        elif melhor_score >= 0.65:
            resultado["perfil_google"] = "Sim"
            resultado["detalhe_google"] = detalhe
        else:
            resultado["perfil_google"] = "Possível"
            resultado["detalhe_google"] = detalhe + " — confira manualmente"
        return resultado
    except requests.RequestException as e:
        return {"perfil_google": "Erro", "detalhe_google": f"Falha de conexão com Google: {e}"}


# ---------------------------------------------------------------------------
# Rotas
# ---------------------------------------------------------------------------
@app.route("/")
def index():
    return render_template("index.html")


@app.route("/api/upload", methods=["POST"])
def upload():
    """Recebe arquivo .xlsx ou .csv/.txt e devolve a lista de CNPJs encontrados."""
    arquivo = request.files.get("arquivo")
    if not arquivo:
        return jsonify({"erro": "Nenhum arquivo enviado"}), 400

    nome = (arquivo.filename or "").lower()
    try:
        if nome.endswith((".xlsx", ".xlsm")):
            from openpyxl import load_workbook

            wb = load_workbook(io.BytesIO(arquivo.read()), read_only=True, data_only=True)
            textos = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    for cel in row:
                        if cel is None:
                            continue
                        if isinstance(cel, (int, float)):
                            # célula numérica: Excel corta zeros à esquerda do CNPJ
                            c = str(int(cel))
                            textos.append(c.zfill(14) if 12 <= len(c) <= 14 else c)
                        else:
                            textos.append(str(cel))
            texto = "\n".join(textos)
        else:
            texto = arquivo.read().decode("utf-8", errors="ignore")
    except Exception as e:
        return jsonify({"erro": f"Não foi possível ler o arquivo: {e}"}), 400

    cnpjs = extrair_cnpjs(texto)
    validos = [c for c in cnpjs if cnpj_valido(c)]
    invalidos = [formatar_cnpj(c) for c in cnpjs if not cnpj_valido(c)]
    return jsonify({"cnpjs": validos, "invalidos": invalidos, "total": len(validos)})


@app.route("/api/consultar", methods=["POST"])
def consultar():
    """Consulta um CNPJ: Receita Federal + Perfil do Google."""
    dados_req = request.get_json(silent=True) or {}
    cnpj = somente_digitos(dados_req.get("cnpj", ""))
    api_key = (dados_req.get("api_key") or "").strip()

    if len(cnpj) != 14:
        return jsonify({"erro": "CNPJ inválido"}), 400

    resultado = {"cnpj": formatar_cnpj(cnpj)}

    receita = consultar_receita(cnpj)
    resultado.update(receita)
    if "erro" in receita:
        resultado["perfil_google"] = "—"
        return jsonify(resultado)

    if api_key:
        resultado.update(consultar_google(api_key, receita))
    else:
        resultado["perfil_google"] = "—"
        resultado["detalhe_google"] = "Sem chave de API do Google"

    return jsonify(resultado)


@app.route("/api/buscar_nome", methods=["POST"])
def buscar_nome():
    """Busca empresas pelo nome fantasia via Google Places (não retorna CNPJ)."""
    dados_req = request.get_json(silent=True) or {}
    nome = (dados_req.get("nome") or "").strip()
    cidade = (dados_req.get("cidade") or "").strip()
    api_key = (dados_req.get("api_key") or "").strip()

    if not nome:
        return jsonify({"erro": "Informe o nome da empresa"}), 400
    if not api_key:
        return jsonify({"erro": "A busca por nome usa o Google Places e exige a chave de API"}), 400

    query = f"{nome} {cidade}".strip()
    headers = {
        "Content-Type": "application/json",
        "X-Goog-Api-Key": api_key,
        "X-Goog-FieldMask": (
            "places.displayName,places.formattedAddress,places.rating,"
            "places.userRatingCount,places.googleMapsUri,places.businessStatus,"
            "places.nationalPhoneNumber,places.websiteUri,"
            "places.types,places.primaryType,places.primaryTypeDisplayName"
        ),
    }
    body = {"textQuery": query, "languageCode": "pt-BR", "regionCode": "BR", "maxResultCount": 10}

    try:
        r = requests.post(PLACES_URL, json=body, headers=headers, timeout=20)
        if r.status_code == 403:
            return jsonify({"erro": "Chave de API inválida ou Places API (New) não habilitada"}), 400
        if r.status_code != 200:
            return jsonify({"erro": f"Erro Google (HTTP {r.status_code})"}), 502

        empresas = []
        for p in r.json().get("places", []):
            grupo = categoria_permitida(p)
            if not grupo:
                continue  # fora das categorias permitidas
            site = p.get("websiteUri", "") or ""
            nome_google = (p.get("displayName") or {}).get("text", "")
            empresas.append({
                "nome_google": nome_google,
                "endereco_google": p.get("formattedAddress", ""),
                "telefone_google": p.get("nationalPhoneNumber", ""),
                "site": site,
                "tem_site": bool(site.strip()),
                "categoria_google": (p.get("primaryTypeDisplayName") or {}).get("text", ""),
                "grupo_categoria": grupo,
                "avaliacao": p.get("rating", ""),
                "total_avaliacoes": p.get("userRatingCount", ""),
                "link_maps": p.get("googleMapsUri", ""),
                "fechado": p.get("businessStatus") == "CLOSED_PERMANENTLY",
                "semelhanca": round(similaridade(nome_google, nome) * 100),
            })

        # alvo da busca: perfis SEM site primeiro, depois por semelhança
        empresas.sort(key=lambda e: (e["tem_site"], -e["semelhanca"]))
        sem_site = sum(1 for e in empresas if not e["tem_site"])
        return jsonify({"empresas": empresas, "total": len(empresas), "sem_site": sem_site})
    except requests.RequestException as e:
        return jsonify({"erro": f"Falha de conexão com Google: {e}"}), 502


if __name__ == "__main__":
    print("\n  Sistema de Busca CNPJ + Perfil Google")
    print("  Acesse: http://localhost:5000\n")
    app.run(host="127.0.0.1", port=5000, debug=False)
